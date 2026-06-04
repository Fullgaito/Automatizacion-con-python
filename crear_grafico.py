import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def find_input_file() -> Path:
    candidates = [
        Path.cwd() / "PRODUCTOS.xlsx",
        Path(__file__).resolve().parent / "PRODUCTOS.xlsx",
    ]
    for path in candidates:
        if path.exists():
            return path
    raise FileNotFoundError(
        "No se encontró 'PRODUCTOS.xlsx'. Coloca el archivo en la carpeta del script o en la carpeta actual."
    )


def prepare_output_folder() -> Path:
    output = Path(__file__).resolve().parent / "graficos_excel"
    output.mkdir(parents=True, exist_ok=True)
    return output


def load_data(path: Path) -> pd.DataFrame:
    df = pd.read_excel(path)

    if "FECHA" not in df.columns:
        raise KeyError("La columna 'FECHA' no está presente en el archivo.")

    df = df.dropna(subset=["FECHA"])
    if df.empty:
        raise ValueError("El archivo no contiene filas válidas con fechas.")

    df["FECHA"] = pd.to_datetime(df["FECHA"], dayfirst=True, errors="coerce")
    df = df.dropna(subset=["FECHA"])
    if df.empty:
        raise ValueError("No se pudieron convertir las fechas en la columna 'FECHA'.")

    df["AÑO"] = df["FECHA"].dt.year
    return df


def build_summary(df: pd.DataFrame) -> pd.DataFrame:
    summary = df.groupby("AÑO", as_index=False).size()
    summary.columns = ["AÑO", "CANTIDAD"]
    return summary.sort_values("AÑO")


def save_excel_with_chart(summary: pd.DataFrame, output_path: Path) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)

    wb = load_workbook(output_path)
    ws = wb["Resumen"]
    max_row = ws.max_row

    chart = BarChart()
    chart.type = "col"
    chart.style = 12
    chart.title = "Cantidad de productos por año"
    chart.x_axis.title = "Año"
    chart.y_axis.title = "Cantidad de productos"
    chart.width = 22
    chart.height = 13
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.legend = None

    data = Reference(ws, min_col=2, min_row=1, max_row=max_row)
    categories = Reference(ws, min_col=1, min_row=2, max_row=max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    ws.add_chart(chart, "E5")

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4F81BD")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=2):
        for cell in row:
            cell.border = border

    wb.save(output_path)


def main() -> None:
    try:
        input_file = find_input_file()
        output_folder = prepare_output_folder()
        output_file = output_folder / "grafico_productos_por_año.xlsx"

        print(f"Cargando datos desde: {input_file}")
        df = load_data(input_file)
        summary = build_summary(df)

        if summary.empty:
            raise ValueError("No se pudo generar el resumen de productos.")

        save_excel_with_chart(summary, output_file)
        print(f"Archivo generado correctamente: {output_file}")
    except Exception as exc:
        print(f"Error: {exc}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
